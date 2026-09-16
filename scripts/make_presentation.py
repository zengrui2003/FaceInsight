from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


BG = RGBColor(18, 24, 38)
ACCENT = RGBColor(80, 227, 194)
ACCENT2 = RGBColor(255, 176, 32)
TEXT = RGBColor(240, 244, 248)
MUTED = RGBColor(181, 193, 205)


def add_slide(prs: Presentation, title: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = BG
    background.line.fill.background()

    accent = slide.shapes.add_shape(1, Inches(0.55), Inches(0.68), Inches(0.18), Inches(0.45))
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT
    accent.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.95), Inches(0.52), Inches(11), Inches(0.8))
    title_frame = title_box.text_frame
    title_frame.text = title
    title_frame.margin_left = 0
    title_frame.word_wrap = True
    paragraph = title_frame.paragraphs[0]
    paragraph.font.size = Pt(30)
    paragraph.font.bold = True
    paragraph.font.color.rgb = TEXT
    return slide


def add_body(slide, lines, top=1.6, left=0.95, width=11.1, height=5.4, size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    text_frame = box.text_frame
    text_frame.word_wrap = True
    for index, line in enumerate(lines):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = TEXT if not line.startswith(("    ", "·")) else MUTED
        paragraph.space_after = Pt(10)
    return box


def add_footer(slide, text):
    box = slide.shapes.add_textbox(Inches(0.95), Inches(6.65), Inches(11), Inches(0.4))
    text_frame = box.text_frame
    text_frame.text = text
    paragraph = text_frame.paragraphs[0]
    paragraph.font.size = Pt(12)
    paragraph.font.color.rgb = MUTED


def build(path: Path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = add_slide(prs, "")
    title_box = slide.shapes.add_textbox(Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.text = "人脸颜值数据洞察"
    title_frame.paragraphs[0].font.size = Pt(54)
    title_frame.paragraphs[0].font.bold = True
    title_frame.paragraphs[0].font.color.rgb = TEXT

    sub_box = slide.shapes.add_textbox(Inches(0.95), Inches(3.55), Inches(11), Inches(1.2))
    sub_frame = sub_box.text_frame
    sub_frame.text = "icrawler · 百度智能云 · SQLite · pyecharts · tkinter"
    sub_frame.paragraphs[0].font.size = Pt(22)
    sub_frame.paragraphs[0].font.color.rgb = ACCENT
    add_footer(slide, "课堂项目拓展答辩")

    slide = add_slide(prs, "项目介绍")
    add_body(
        slide,
        [
            "一句话介绍：从图片抓取到人脸分析，再到历史数据洞察的一体化桌面工具。",
            "输入关键词，自动用 icrawler 从 Bing 抓取图片。",
            "调用百度智能云人脸检测，解析 age、expression、beauty 字段。",
            "把每个关键词的检测统计写入 SQLite，形成可追溯数据资产。",
            "用 pyecharts 生成分数柱状图与历史数据大屏。",
        ],
    )
    add_footer(slide, "课堂内容保留：爬虫 + 人脸检测 + 图表")

    slide = add_slide(prs, "功能亮点")
    add_body(
        slide,
        [
            "· tkinter 桌面界面，后台线程执行，进度条和运行日志实时刷新。",
            "· 本地图片文件夹分析模式，爬虫不可用时仍能演示核心能力。",
            "· 百度返回异常分类处理：无人脸、图片超限、密钥或权限问题给出提示。",
            "· 修复了课堂笔记中的除零 bug：全无人脸时平均分保存为 NULL。",
            "· 历史记录表、CSV 导出、最新图表自动打开。",
            "· 大屏新增颜值仪表盘、分数雷达、环形占比、渐变面积趋势图。",
        ],
        size=20,
    )
    add_footer(slide, "稳定演示优先：现场断网也有本地兜底")

    slide = add_slide(prs, "技术架构")
    add_body(
        slide,
        [
            "数据采集层：icrawler BingImageCrawler，多线程下载图片。",
            "算法服务层：baidu-aip AipFace.detect，Base64 上传，返回 JSON。",
            "数据存储层：SQLite beauty_analysis 表，建表、插入、查询复用课堂知识。",
            "可视化层：pyecharts Bar / Line / Pie / Page，导出交互式 HTML。",
            "表现层：tkinter 主界面 + threading 后台任务 + queue 更新 UI。",
        ],
        size=20,
    )
    add_footer(slide, "模块边界清晰，便于后续扩展")

    slide = add_slide(prs, "关键代码：数据采集")
    add_body(
        slide,
        [
            "folder = root_dir / safe_folder_name(keyword)",
            "if folder.exists(): shutil.rmtree(folder)",
            "folder.mkdir(parents=True, exist_ok=True)",
            "crawler = BingImageCrawler(",
            "    feeder_threads=2, parser_threads=4, downloader_threads=8,",
            "    storage={'root_dir': str(folder)})",
            "crawler.crawl(keyword=keyword, max_num=max_num)",
            "    ",
            "亮点：每次分析前清理旧图，避免历史缓存污染统计结果。",
        ],
        size=18,
    )
    add_footer(slide, "face_insight/crawler.py")

    slide = add_slide(prs, "关键代码：百度人脸检测")
    add_body(
        slide,
        [
            "content = base64.b64encode(path.read_bytes()).decode('utf-8')",
            "options = {'face_field': 'age,expression,beauty'}",
            "response = self.client.detect(content, 'BASE64', options)",
            "faces = response.get('result', {}).get('face_list', [])",
            "beauty = float(faces[0].get('beauty', 0))",
            "return min(10, max(1, int(beauty // 10) + 1))",
            "    ",
            "亮点：0 表示无人脸，1-10 表示颜值分段；错误码单独处理，不会直接崩掉界面。",
        ],
        size=18,
    )
    add_footer(slide, "face_insight/face_analysis.py")

    slide = add_slide(prs, "关键代码：数据库与除零修复")
    add_body(
        slide,
        [
            "face_count = sum(score_counts[1:])",
            "if face_count == 0:",
            "    return None",
            "total = sum(score * count for score, count in enumerate(score_counts))",
            "return round(total / face_count, 2)",
            "    ",
            "表设计：keyword、count_has_face、count_no_face、num_one-num_ten、beauty_avg。",
            "全部使用参数化 SQL，避免拼接字符串带来的风险。",
        ],
        size=18,
    )
    add_footer(slide, "face_insight/storage.py")

    slide = add_slide(prs, "关键代码：可视化")
    add_body(
        slide,
        [
            "Bar().add_xaxis(SCORE_LABELS).add_yaxis('图片数量', items)",
            "Pie().add('占比', data, radius=['38%', '68%'])",
            "Gauge().add('颜值指数', [('得分', value)], min_=0, max_=10)",
            "Page(layout=Page.DraggablePageLayout)",
            "page.add(build_score_chart(...))",
            "page.add(build_average_trend(records))",
            "page.add(build_face_distribution(records))",
            "page.add(build_keyword_comparison(records))",
            "page.render('exports/人脸检测历史数据大屏.html')",
        ],
        size=18,
    )
    add_footer(slide, "face_insight/charts.py")

    slide = add_slide(prs, "现场演示")
    add_body(
        slide,
        [
            "1. 输入关键词，例如“学生证件照”，点击开始分析。",
            "2. 观察爬取、百度检测、数据库保存、图表生成四个阶段。",
            "3. 打开最新柱状图，说明无人脸和 1-10 分分布。",
            "4. 连续分析两个关键词后打开历史大屏。",
            "5. 展示颜值仪表盘、趋势图、雷达图、分布占比和 CSV 导出。",
            "备注：建议录制 60-90 秒演示视频作为备份。",
        ],
        size=20,
    )
    add_footer(slide, "演示脚本详见 docs/演示讲稿.md")

    slide = add_slide(prs, "问题排查")
    add_body(
        slide,
        [
            "· 有人脸数量为 0 时除法报错：改为返回 NULL，并在图表中显示“无有效人脸”。",
            "· 爬虫下载的不是有效图片：只分析 jpg/jpeg/png/bmp，其他格式跳过。",
            "· 百度返回无人脸：222202、222203、222204 等错误码归类为 score=0。",
            "· 图片超过 10MB：提前拦截，并提示文件名和限制。",
            "· 爬虫受网络或站点策略影响：增加本地文件夹分析模式。",
            "· GUI 卡死：爬取和检测放入子线程，Tkinter 主线程只负责渲染。",
        ],
        size=20,
    )
    add_footer(slide, "每个问题都有定位过程、处理方案和验证方式")

    slide = add_slide(prs, "项目价值与延展")
    add_body(
        slide,
        [
            "课堂项目：完成了一组工具库的串联调用。",
            "拓展项目：变成有界面、有状态、有兜底、有历史、可复盘的小型数据产品。",
            "后续可延展：增加用户登录与权限、人脸分组统计、定时任务、多人协作看板。",
            "也可以把爬虫替换为公开数据集，把人脸服务替换为其他 AI 能力。",
        ],
        size=20,
    )
    add_footer(slide, "从作业到产品思维")

    slide = add_slide(prs, "感谢致辞")
    add_body(
        slide,
        [
            "感谢老师课堂讲解和项目指导，让我们从零走到了完整作品。",
            "感谢小组成员分工协作，一起排查了爬虫、接口和数据库问题。",
            "感谢百度智能云与开源社区提供的文档和示例。",
            "我们也会继续保持对数据和 AI 的好奇心，把这些方法用到更有价值的问题上。",
            "欢迎各位老师同学提问！",
        ],
        size=22,
    )
    add_footer(slide, "Thank You")

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


if __name__ == "__main__":
    build(Path("PPT/人脸颜值数据洞察-答辩.pptx"))
